"""Create a separate organized list; preserve all distinct roles and names."""
import json, re
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
data = json.loads((root / 'names.js').read_text(encoding='utf8').split(' = ', 1)[1].rstrip(';\n'))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(19)
sec.left_margin = sec.right_margin = Mm(20)
for name, size in [('Normal',12),('Title',20),('Heading 1',16),('Heading 2',13)]:
    st = doc.styles[name]
    st.font.name = 'Leelawadee UI'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    rf = st.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii','hAnsi','eastAsia','cs']:
        rf.set(qn('w:'+key), 'Leelawadee UI')
    st.paragraph_format.space_after = Pt(1)
    st.paragraph_format.line_spacing = 1.1
doc.add_paragraph('รายชื่อผู้ร่วมทำบุญกฐินสามัคคี', 'Title')
doc.add_paragraph('วัดเขาพนมโดมวนาราม ตำบลวัดตะแบก อำเภอเทพสถิต จังหวัดชัยภูมิ')
doc.add_paragraph('รายชื่อผู้ร่วมทำบุญ แบ่งตามสายและตำแหน่ง รวม 126 รายการ')
total = 0
removed = []
for gi, group in enumerate(data['groups']):
    if gi:
        doc.add_page_break()
    doc.add_paragraph(group['title'], 'Heading 1')
    for si, section in enumerate(group['sections']):
        if gi == 0 and si >= 2:
            doc.add_page_break()
        heading = doc.add_paragraph(section['title'], 'Heading 2')
        if gi == 2:
            heading.paragraph_format.space_before = Pt(6)
        seen = set()
        ordinal = 0
        for entry in section['entries']:
            text = re.sub(r'^\d+\.\s*', '', entry['text'])
            key = re.sub(r'\s+', ' ', text).strip()
            if key in seen:
                removed.append((group['title'],section['title'],text))
                continue
            seen.add(key)
            ordinal += 1
            p = doc.add_paragraph(f'{ordinal}. {text}')
            p.paragraph_format.left_indent = Mm(9)
            p.paragraph_format.first_line_indent = Mm(-9)
            p.paragraph_format.keep_together = True
            if gi == 2:
                p.runs[0].font.size = Pt(11)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
            total += 1
assert total == 126
out = root / 'output'
out.mkdir(exist_ok=True)
target = out / 'รายชื่อกฐินสามัคคี_ฉบับจัดระเบียบ.docx'
for element in doc.styles.element.iter():
    for child in list(element):
        if child.tag == qn('w:pBdr'):
            element.remove(child)
doc.save(target)
verify = Document(target)
actual = [p.text for p in verify.paragraphs if re.match(r'^\d+\.',p.text)]
assert len(actual) == total
assert '20. สท.รุ่งอรุณ หายโศก และครอบครัว' in actual
print(json.dumps({'output':str(target),'entries':total,'within_section_duplicates':removed},ensure_ascii=True))
